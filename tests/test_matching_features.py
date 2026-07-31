import io

from docx import Document

import config
from backend.cv_service import extract_text
from backend.models import AnalysisResult
from backend.profile_service import (
    build_analysis_context,
    build_profile_snapshot,
    profile_for_search,
)
from backend.repository import (
    create_cover_letter,
    create_run,
    get_or_create_search_settings,
    latest_cover_letter,
    list_jobs,
    search_settings_dict,
    set_job_dismissed,
    update_cover_letter_content,
    update_search_settings,
    upsert_job,
)
from backend.run_manager import RunManager
from backend.schemas import RunType


def test_search_controls_persist(session):
    initial = get_or_create_search_settings(session)
    values = search_settings_dict(initial)
    values.update(
        {
            "keywords": ["Security Engineer", "DevOps Engineer"],
            "minimum_match_score": 75,
            "number_of_jobs": 12,
            "sources": ["LinkedIn", "Greenhouse", "Lever"],
            "greenhouse_boards": ["Example AG | example"],
            "lever_sites": ["Example AG | example"],
        }
    )
    updated = update_search_settings(session, values)
    session.commit()

    persisted = search_settings_dict(updated)
    assert persisted["keywords"] == ["Security Engineer", "DevOps Engineer"]
    assert persisted["minimum_match_score"] == 75
    assert persisted["number_of_jobs"] == 12
    assert persisted["greenhouse_boards"] == ["Example AG | example"]
    assert persisted["lever_sites"] == ["Example AG | example"]


def test_search_is_the_single_public_pipeline_endpoint():
    from backend.app import app

    paths = {route.path for route in app.routes}
    assert RunType.SEARCH.value == "search"
    assert "/api/runs/search" in paths
    assert "/api/runs/scraper" not in paths
    assert "/api/runs/analyzer" not in paths
    assert "/api/runs/full" not in paths


def test_search_run_discovers_then_scores(monkeypatch):
    manager = RunManager()
    calls: list[tuple[str, object]] = []
    manager._active_run_id = "search-run"

    monkeypatch.setattr(manager, "_update_run", lambda *_args, **_kwargs: None)
    monkeypatch.setattr(
        manager,
        "_scrape",
        lambda _run_id, _parameters: calls.append(("discover", None)) or [11, 12],
    )
    monkeypatch.setattr(
        manager,
        "_analyze",
        lambda _run_id, *, job_ids, force: calls.append(("score", (job_ids, force))),
    )

    manager._execute("search-run", "search", {"search_controls": {}})

    assert calls == [("discover", None), ("score", ([11, 12], False))]


def test_analysis_snapshot_hashes_but_does_not_expose_raw_cv():
    profile = {"full_name": "Ada Candidate", "skills": ["Python"]}
    settings = {"minimum_match_score": 60, "keywords": ["Security Engineer"]}
    raw_cv = "Private CV text with a confidential phone number 555-0100."

    snapshot = build_profile_snapshot(profile, settings, raw_cv)
    context = build_analysis_context(profile, settings, raw_cv)

    assert raw_cv not in snapshot
    assert "SHA-256" in snapshot
    assert raw_cv in context


def test_search_keywords_are_the_single_role_target_source():
    profile = {"role_keywords": ["Old profile role"], "skills": ["Python"]}
    settings = {"keywords": ["Security Engineer", "DevOps Engineer"]}

    search_profile = profile_for_search(profile, settings)

    assert search_profile["role_keywords"] == settings["keywords"]
    assert profile["role_keywords"] == ["Old profile role"]


def test_scored_jobs_are_thresholded_sorted_and_dismissible(session):
    run = create_run(session, "score-run", "analyzer", {})
    low_job, _ = upsert_job(
        session,
        {
            "title": "Support Specialist",
            "company": "Example",
            "link": "https://linkedin.com/jobs/view/6612345678",
            "description": "General support.",
        },
    )
    high_job, _ = upsert_job(
        session,
        {
            "title": "Security Engineer",
            "company": "Example",
            "link": "https://linkedin.com/jobs/view/6612345679",
            "description": "Cloud security engineering.",
        },
    )
    session.add_all(
        [
            AnalysisResult(
                job_id=low_job.id,
                run_id=run.id,
                candidate_profile_snapshot="profile",
                candidate_profile_version=1,
                is_good_match=False,
                seniority_ok=True,
                verdict="Low alignment.",
                gemini_model=config.GEMINI_MODEL,
                match_score=58,
                qualifies=False,
                recommendation_label="Low match",
            ),
            AnalysisResult(
                job_id=high_job.id,
                run_id=run.id,
                candidate_profile_snapshot="profile",
                candidate_profile_version=1,
                is_good_match=True,
                seniority_ok=True,
                verdict="Strong alignment.",
                gemini_model=config.GEMINI_MODEL,
                match_score=87,
                qualifies=True,
                recommendation_label="Strong match",
            ),
        ]
    )
    session.commit()

    rows, total = list_jobs(
        session,
        minimum_score=60,
        sort="match_quality",
        direction="desc",
    )
    assert total == 1
    assert rows[0][0].id == high_job.id
    assert rows[0][1].match_score == 87

    set_job_dismissed(session, high_job, True)
    session.commit()
    rows, total = list_jobs(session, minimum_score=60)
    assert rows == []
    assert total == 0


def test_hard_match_filters_are_deterministic():
    evaluation = {
        "match_score": 82,
        "seniority_ok": True,
        "critical_gaps": [],
        "required_experience_years": 2,
        "work_model": "Hybrid",
        "required_languages": ["English"],
        "job_location": "Zurich, Switzerland",
    }
    profile = {"languages": [{"language": "English", "proficiency": "C1"}]}
    settings = {
        "minimum_match_score": 75,
        "include_stretch_roles": False,
        "max_required_experience_years": 4,
        "work_models": ["Hybrid"],
        "exclude_unavailable_languages": True,
        "target_locations": ["Zurich, Switzerland"],
        "exclude_outside_locations": True,
    }
    assert RunManager._qualifies(evaluation, profile, settings) is True

    evaluation["required_experience_years"] = 6
    assert RunManager._qualifies(evaluation, profile, settings) is False


def test_docx_text_extraction_and_editable_cover_letter(session):
    document = Document()
    document.add_heading("Ada Candidate", level=1)
    document.add_paragraph("Security engineering experience with Python and Azure.")
    buffer = io.BytesIO()
    document.save(buffer)

    raw_text = extract_text(buffer.getvalue(), ".docx")
    assert "Ada Candidate" in raw_text
    assert "Python and Azure" in raw_text

    job, _ = upsert_job(
        session,
        {
            "title": "Security Engineer",
            "company": "Example",
            "link": "https://linkedin.com/jobs/view/7712345678",
            "description": "Security role.",
        },
    )
    letter = create_cover_letter(
        session,
        job_id=job.id,
        analysis_result_id=None,
        profile_version=2,
        content="Dear Hiring Team,\n\nInitial tailored letter content.",
        model=config.GEMINI_MODEL,
    )
    session.commit()
    update_cover_letter_content(
        session,
        letter,
        "Dear Hiring Team,\n\nEdited truthful tailored letter content.",
    )
    session.commit()

    stored = latest_cover_letter(session, job.id)
    assert stored is not None
    assert stored.id == letter.id
    assert "Edited truthful" in stored.content
