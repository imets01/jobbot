"""Private CV storage, text extraction, and structured Gemini extraction."""

from __future__ import annotations

import io
import json
import re
from pathlib import Path
from typing import Any
from uuid import uuid4

from docx import Document
from google.genai import types
from pypdf import PdfReader

import config
from analyzer import _parse_structured_response, build_client


ALLOWED_EXTENSIONS = {".pdf", ".docx"}
ALLOWED_CONTENT_TYPES = {
    "application/pdf",
    "application/octet-stream",
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
}

EXTRACTION_SCHEMA = {
    "type": "object",
    "properties": {
        "full_name": {"type": "string"},
        "email": {"type": "string"},
        "current_location": {"type": "string"},
        "linkedin_url": {"type": "string"},
        "portfolio_url": {"type": "string"},
        "education": {"type": "array", "items": {"type": "string"}},
        "work_experience": {"type": "array", "items": {"type": "string"}},
        "skills": {"type": "array", "items": {"type": "string"}},
        "certifications": {"type": "array", "items": {"type": "string"}},
        "languages": {
            "type": "array",
            "items": {
                "type": "object",
                "properties": {
                    "language": {"type": "string"},
                    "proficiency": {"type": "string"},
                },
                "required": ["language", "proficiency"],
            },
        },
        "projects": {"type": "array", "items": {"type": "string"}},
        "professional_summary": {"type": "string"},
    },
    "required": [
        "full_name",
        "email",
        "current_location",
        "linkedin_url",
        "portfolio_url",
        "education",
        "work_experience",
        "skills",
        "certifications",
        "languages",
        "projects",
        "professional_summary",
    ],
}


def validate_cv_upload(filename: str, content_type: str | None, size: int) -> str:
    """Validate a local CV upload and return its normalized extension."""
    extension = Path(filename).suffix.casefold()
    if extension not in ALLOWED_EXTENSIONS:
        raise ValueError("CV must be a PDF or DOCX file.")
    if content_type and content_type not in ALLOWED_CONTENT_TYPES:
        raise ValueError("CV content type must be PDF or DOCX.")
    if size <= 0:
        raise ValueError("CV file is empty.")
    if size > config.MAX_CV_SIZE_BYTES:
        raise ValueError("CV file exceeds the 10 MB limit.")
    return extension


def store_cv_bytes(filename: str, extension: str, content: bytes) -> Path:
    """Store a CV under an opaque local filename."""
    safe_stem = re.sub(r"[^a-zA-Z0-9_-]+", "-", Path(filename).stem).strip("-")[:80]
    path = config.ensure_cv_dir() / f"{safe_stem or 'cv'}-{uuid4().hex[:12]}{extension}"
    path.write_bytes(content)
    return path


def extract_text(content: bytes, extension: str) -> str:
    """Extract text locally from a supported PDF or DOCX document."""
    if extension == ".pdf":
        reader = PdfReader(io.BytesIO(content))
        return "\n\n".join((page.extract_text() or "").strip() for page in reader.pages).strip()
    document = Document(io.BytesIO(content))
    paragraphs = [paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip()]
    for table in document.tables:
        for row in table.rows:
            values = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if values:
                paragraphs.append(" | ".join(values))
    return "\n".join(paragraphs).strip()


def extract_structured_profile(raw_text: str) -> dict[str, Any]:
    """Use Gemini to convert raw CV text into explicitly supported profile fields."""
    if len(raw_text.strip()) < 20:
        raise ValueError("No usable text could be extracted from the CV.")
    prompt = f"""Extract factual candidate data from the CV below.
Use only information explicitly present in the CV. Use empty strings or arrays when unknown.
Keep work experience, education, projects, certifications, and language entries concise but informative.
Write a two-to-four sentence professional summary with no invented claims.

CV text:
```
{raw_text[:40_000]}
```
"""
    client = build_client()
    response = client.models.generate_content(
        model=config.GEMINI_MODEL,
        contents=prompt,
        config=types.GenerateContentConfig(
            system_instruction=(
                "You are a precise CV parser. Extract only supported facts and never infer "
                "credentials, dates, proficiency, or experience that are not written."
            ),
            response_mime_type="application/json",
            response_schema=EXTRACTION_SCHEMA,
            thinking_config=types.ThinkingConfig(thinking_budget=0),
            max_output_tokens=4096,
        ),
    )
    result = _parse_structured_response(response.text or "{}")
    return {
        key: result.get(key, [] if EXTRACTION_SCHEMA["properties"][key]["type"] == "array" else "")
        for key in EXTRACTION_SCHEMA["properties"]
    }


def delete_stored_cv(storage_path: str | None) -> None:
    """Delete a replaced local CV without allowing arbitrary path deletion."""
    if not storage_path:
        return
    candidate = Path(storage_path).resolve()
    root = config.ensure_cv_dir().resolve()
    if root in candidate.parents and candidate.is_file():
        candidate.unlink(missing_ok=True)
